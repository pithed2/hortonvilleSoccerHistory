from pathlib import Path
import math
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT=Path(__file__).parent
font='C:/Windows/Fonts/arial.ttf'
def diagram(kind):
 im=Image.new('RGB',(1500,630),'white'); d=ImageDraw.Draw(im)
 def txt(x,y,t,size=25): d.text((x,y),t,font=ImageFont.truetype(font,size),fill='#17252b')
 def line(a,b,color='#273c45',dash=False,width=5):
  if dash:
   length=math.dist(a,b)
   for i in range(0,int(length),24):
    p=i/length; q=min((i+13)/length,1);d.line((a[0]+(b[0]-a[0])*p,a[1]+(b[1]-a[1])*p,a[0]+(b[0]-a[0])*q,a[1]+(b[1]-a[1])*q),fill=color,width=width)
  else:d.line((*a,*b),fill=color,width=width)
  an=math.atan2(b[1]-a[1],b[0]-a[0]); pts=[b,(b[0]-20*math.cos(an-.5),b[1]-20*math.sin(an-.5)),(b[0]-20*math.cos(an+.5),b[1]-20*math.sin(an+.5))];d.polygon(pts,fill=color)
 def player(x,y,t='A',c='#176da0'):
  d.ellipse((x-21,y-21,x+21,y+21),fill=c,outline='white',width=2);d.text((x-9,y-15),t,font=ImageFont.truetype(font,25),fill='white')
 def ball(x,y):d.ellipse((x-8,y-8,x+8,y+8),fill='black')
 d.rectangle((100,65,1370,495),fill='#f0f6ee',outline='#42634a',width=5)
 if kind==0:
  d.line((735,65,735,495),fill='#42634a',width=4)
  txt(570,15,'20 yd long x 10 yd wide')
  for x,y in [(330,180),(400,380)]:player(x,y)
  for x,y in [(1080,180),(1130,380)]:player(x,y,'D','#ba493b')
  player(560,410,'N','#8b6a12');ball(370,185);line((380,180),(1030,180));txt(610,255,'Low net or cone line')
  txt(130,515,'One court shown   |   Neutral joins the receiving side   |   Fifth court plays 3v3',26)
 else:
  for x in [75,1370]:d.rectangle((x,235,x+25,325),outline='#273c45',width=5)
  if kind==1:
   txt(580,15,'30 yd long x 20 yd wide');d.line((1120,65,1120,495),fill='#91a58c',width=3);d.line((350,65,350,495),fill='#91a58c',width=3)
   for x,y in [(340,280),(600,150),(650,420),(880,370),(920,180)]:player(x,y)
   for x,y in [(480,340),(760,310),(1010,150),(1010,320),(630,100)]:player(x,y,'D','#ba493b')
   player(550,300,'N','#8b6a12');ball(375,280)
   line((385,280),(1170,220));line((920,180),(1190,220),dash=True);line((1200,230),(1340,280),color='#ba493b');txt(1135,80,'5 yd',24)
   txt(130,515,'Example   Pass into the end zone as the runner breaks beyond the defender',26)
  elif kind==2:
   txt(570,15,'20 yd long x 20 yd wide')
   for x,y,t in [(330,290,'1'),(720,140,'2'),(680,380,'3'),(450,440,'4')]:player(x,y,t)
   ball(370,290);line((380,290),(550,270));txt(390,220,'Dribble',24)
   line((570,270),(1090,190));line((745,140),(1110,190),dash=True);line((700,380),(1030,340),dash=True);line((470,430),(700,300),dash=True)
   line((1120,195),(1340,280),color='#ba493b');txt(950,400,'Next repetition attacks left',24)
   txt(130,515,'One group shown   |   Player 1 releases the runner   |   Rotate passer and finishers',26)
  else:
   txt(560,15,'Half field   |   Use available width')
   d.line((735,65,735,495),fill='#91a58c',width=3)
   for x,y in [(220,280),(370,110),(370,280),(370,445),(570,130),(570,300),(570,450),(850,120),(860,290),(860,430),(1030,220)]:player(x,y)
   for x,y in [(1250,280),(1130,100),(1130,290),(1130,440),(980,110),(980,350),(980,455),(770,195),(770,380),(650,220),(650,400)]:player(x,y,'D','#ba493b')
   player(720,100,'N','#8b6a12');ball(895,290);line((900,280),(1210,180));line((1045,220),(1210,175),dash=True);line((1230,180),(1340,270),color='#ba493b')
   txt(130,515,'11v11 plus neutral   |   Example attacking pattern shown   |   Reverse on turnovers',26)
 txt(130,570,'A  Attackers     D  Defenders     N  Neutral     Solid arrow  Ball     Dashed arrow  Run',25)
 im.save(OUT/f'diagram-{kind}.png')

doc=Document(); sec=doc.sections[0];sec.top_margin=sec.bottom_margin=Inches(.65);sec.left_margin=sec.right_margin=Inches(.7)
for name in ['Normal','Title','Heading 1','Heading 2']:
 s=doc.styles[name];s.font.name='Arial';s.font.color.rgb=RGBColor(0,0,0)
doc.styles['Normal'].font.size=Pt(10.5);doc.styles['Normal'].paragraph_format.space_after=Pt(6)
doc.styles['Title'].font.size=Pt(25);doc.styles['Heading 1'].font.size=Pt(19);doc.styles['Heading 2'].font.size=Pt(12)
def p(t):doc.add_paragraph(t)
def h(t):doc.add_heading(t,2)
def bullets(items):
 for t in items:doc.add_paragraph(t,'List Bullet')
doc.add_paragraph('Through Balls and Final Third Finishing','Title')
p('95 minute soccer session plan | 23 players | Coach __________________')
p('Develop the timing of forward runs, the weight of passes behind defenders, and composed finishing. Progress from ball control to transition play, finishing combinations, and a conditioned game.')
h('Session schedule')
table=doc.add_table(rows=1, cols=3);table.style='Table Grid'
for c,t in zip(table.rows[0].cells,['Time','Activity','Duration']):c.text=t
for row in [('0–10','Warmup','10 min'),('10–20','Arrival soccer tennis','10 min'),('20–40','Through balls in transition','20 min'),('40–60','Final third finishing','20 min'),('60–90','Attacking game','30 min'),('90–95','Cooldown and review','5 min')]:
 for c,t in zip(table.add_row().cells,row):c.text=t
for row in table.rows:
 for c in row.cells:
  pr=c._tc.get_or_add_tcPr();b=OxmlElement('w:tcBorders')
  for side in ['top','left','bottom','right']:
   e=OxmlElement('w:'+side);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');b.append(e)
  pr.append(b)
h('Equipment and organization')
p('Bring at least 12 balls, 40–50 cones, two sets of 11 bibs plus a neutral bib, and 2–4 mini goals. Cone gates replace extra goals during station work. Set up stations before arrival and include rotations and water within activity time.')
h('Warmup from minute 0 to 10')
p('Complete your usual 10 minute team warmup before soccer tennis. Build from movement and mobility into ball work and progressive accelerations.')
h('Coaching language')
bullets(['Look up before the pass. Release when the runner can arrive first.','Curve or delay the run; accelerate as the passer looks up.','Pass into the runner’s path. Take the first touch toward the finish.'])
h('Cooldown and review from minute 90 to 95')
p('Two minutes of easy walking and gentle mobility, then three minutes to review: What tells you the through ball is on? How did your run help the passer? Ask players to identify one successful example.')

activities=[
('Arrival soccer tennis','10–20 minutes | 10 minutes',0,
 'Set five courts, each 20 x 10 yards, with a center cone line or low net. Four courts play 2v2 and one plays 3v3, using 22 players. Player 23 is a roaming neutral who joins the receiving side. Use one ball per court; mini goals are optional targets for a variation.',
 ['Start with a controlled serve across the middle. The receiving side combines and returns the ball inside the opposite half.','Allow one bounce and up to three team touches initially. A second bounce or a ball out gives the other side a point.','Rotate opponents and the neutral every two minutes. Keep rallies moving with spare balls.'],
 ['Communicate early and adjust your feet before contact.','Use a controlled first touch and an accurate return. Support your partner at an angle.'],
 'Easier: permit two bounces or a catch to restart a broken rally. Harder: remove the bounce or limit players to one or two touches. Optional goal version: award a bonus for a controlled return through a target gate.'),
('Through balls in transition','20–40 minutes | 20 minutes',1,
 'Use two 30 x 20 yard grids, each with a mini goal or cone gate at both ends and a 5 yard end zone. Grid A plays 5v5 plus a neutral; Grid B plays 6v6. The neutral supports possession. Widen a grid if congestion prevents passing lanes.',
 ['Play four five-minute rounds, including short resets. On a turnover, attack the opposite goal immediately.','A goal counts only after a forward pass into the attacking end zone reaches a teammate running in from outside it. Players cannot wait in the end zone.','After receiving, finish promptly. Defenders can recover into the zone. Restart from the conceding team’s end after a goal; rotate the neutral and matchups between rounds.'],
 ['Scan for a forward lane immediately after winning possession.','Pass ahead of the runner with enough pace to beat the defender.','Delay the run until the passer can release; stay wide enough to open a lane.'],
 'Easier: widen the grid or allow a free first touch after a regain. Harder: use two touches in midfield while keeping the receiver free to finish.'),
('Final third finishing','40–60 minutes | 20 minutes',2,
 'Set five 20 x 20 yard grids with goals or cone gates on opposite sides. Use groups of 4, 4, 5, 5 and 5. Four players work per repetition; in groups of five, the spare player serves or retrieves and rotates in after each attempt. Use one active ball and a spare per grid.',
 ['The ball carrier dribbles forward. Three teammates make different supporting runs: one beyond, one across and one underneath.','The carrier passes to a forward runner, who finishes after receiving. No direct goal by the original dribbler.','One attack per grid at a time. Retrieve the ball and attack the opposite end on the next repetition. Rotate the passer every attempt; change starting positions every five minutes.'],
 ['Create separation before accelerating into the passing lane.','Keep the pass in front of the runner; avoid forcing the receiver to stop.','Look at the target before striking. Use a settling touch when needed.'],
 'Easier: allow a touch to set the ball before shooting. Harder: require a one-touch finish. In groups of five, the extra player can become a recovering defender once the passing pattern is secure.'),
('Attacking game with through balls','60–90 minutes | 30 minutes',3,
 'Use half a field with one mini goal at each end, two goals total. Play 11v11 plus a neutral who supports possession. Use the available width; widen or lengthen if the pitch is too crowded. Play without goalkeepers and do not allow players to stand in the goal mouth.',
 ['Play three ten-minute blocks, including brief water and feedback. Rotate the neutral at each break.','First block: goals count only after a through ball behind the defensive line. Apply offside in the attacking half, judged by the coach.','Second block: all goals count one; a goal from a through ball counts two. Final block: remove scoring conditions and observe whether players choose the pass at the right moment.','On a turnover, the neutral immediately supports the new team. Restart quickly from the end line after goals and with a pass-in from touchlines.'],
 ['Create width and depth so the ball carrier has a forward option.','Recognize when to play through and when to keep possession.','Support the runner for a second pass, rebound or finish. Praise well-timed runs even when the final pass misses.'],
 'Easier: give the neutral unlimited touches and encourage early forward support. Harder: limit midfield players to two touches. With 22 players, remove the neutral and play 11v11.')]
for title,time,k,setup,steps,cues,var in activities:
 doc.add_page_break();doc.add_heading(title,1);p(time)
 diagram(k);doc.add_picture(str(OUT/f'diagram-{k}.png'),width=Inches(7.05))
 doc.inline_shapes[-1]._inline.docPr.set('descr',title+' field diagram. Solid arrows show ball movement and dashed arrows show player runs. '+setup)
 h('Setup');p(setup);h('How to play');bullets(steps);h('Coaching points');bullets(cues);h('Variations');p(var)
for el in list(doc.styles.element.iter(qn('w:pBdr'))): el.getparent().remove(el)
for el in list(doc.element.iter(qn('w:pBdr'))): el.getparent().remove(el)
doc.save(OUT/'Soccer Session Plan.docx')
print(OUT/'Soccer Session Plan.docx')


